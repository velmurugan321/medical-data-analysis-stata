from fastapi import FastAPI, UploadFile, File, HTTPException, Form
from fastapi.middleware.cors import CORSMiddleware
import io, json
import numpy as np
import pandas as pd
from fastapi.responses import StreamingResponse
from docx import Document
from .analysis_engine import descriptive, categorical_association, diagnostic_accuracy, roc_auc, robust_poisson, logistic_regression, table_one, data_quality
from .openeepi_engine import screening as openeepi_screening, calculate as openeepi_calculate
from .dummy_table_engine import analyse_dummy_table
from .dummy_table_parser import build_spec
from .dummy_table_fill import fill_dummy_table
from .rr_analysis import crude_rr, adjusted_rr, format_rr, _binary

app = FastAPI(title='Medical Data Analysis API', version='1.2.1')
ALLOWED_SUFFIXES={'.csv','.xlsx','.xls','.dta','.tsv','.docx'}
app.add_middleware(CORSMiddleware,allow_origins=['*'],allow_credentials=False,allow_methods=['*'],allow_headers=['*'])

def load_dataframe(filename,content):
    name=filename.lower(); suffix=next((s for s in ALLOWED_SUFFIXES if name.endswith(s)),None)
    if suffix is None: raise HTTPException(400,'Unsupported file type. Use CSV, XLSX, XLS, DTA, TSV or DOCX.')
    try:
        if suffix=='.csv': return pd.read_csv(io.BytesIO(content))
        if suffix=='.tsv': return pd.read_csv(io.BytesIO(content),sep='\t')
        if suffix in {'.xlsx','.xls'}: return pd.read_excel(io.BytesIO(content))
        if suffix=='.dta': return pd.read_stata(io.BytesIO(content))
        if suffix=='.docx':
            doc=Document(io.BytesIO(content))
            if not doc.tables: raise ValueError('DOCX contains no tables. Upload a DOCX with a structured data table.')
            tables=[]
            for table in doc.tables:
                rows=[[cell.text.strip() for cell in row.cells] for row in table.rows]
                if rows: tables.append(rows)
            if not tables: raise ValueError('DOCX contains no readable table rows.')
            rows=max(tables,key=lambda x: len(x)*max(len(r) for r in x))
            width=max(len(r) for r in rows); rows=[r+['']*(width-len(r)) for r in rows]
            header=rows[0]
            if len(set(header)) != len(header): header=[f'V{i+1}' for i in range(width)]
            return pd.DataFrame(rows[1:],columns=header)
    except HTTPException: raise
    except Exception as exc: raise HTTPException(422,f'Unable to read dataset: {exc}') from exc

def variable_metadata(df):
    rows=[]
    for col in df.columns:
        s=df[col]; numeric=pd.api.types.is_numeric_dtype(s)
        rows.append({'name':str(col),'type':'numeric' if numeric else 'categorical','storage_type':str(s.dtype),'n':int(s.notna().sum()),'missing':int(s.isna().sum()),'unique':int(s.nunique(dropna=True)),'min':float(s.min()) if numeric and s.notna().any() else None,'max':float(s.max()) if numeric and s.notna().any() else None})
    return rows

@app.get('/health')
def health(): return {'status':'ok','service':'medical-data-analysis-api','version':app.version}

@app.post('/api/v1/data/inspect')
async def inspect_data(file:UploadFile=File(...)):
    df=load_dataframe(file.filename or 'upload.csv',await file.read()); return {'filename':file.filename,'rows':int(df.shape[0]),'columns':int(df.shape[1]),'variables':variable_metadata(df)}

@app.post('/api/v1/data/document')
async def inspect_document(file:UploadFile=File(...)):
    name=file.filename or 'document.docx'; content=await file.read()
    if not name.lower().endswith('.docx'): raise HTTPException(400,'Document extraction currently supports DOCX. DOC/PDF/RTF remain reference uploads.')
    try:
        doc=Document(io.BytesIO(content)); paragraphs=[p.text.strip() for p in doc.paragraphs if p.text.strip()]; tables=[]
        for i,table in enumerate(doc.tables,1):
            rows=[[cell.text.strip() for cell in row.cells] for row in table.rows]
            if rows: tables.append({'table_number':i,'rows':rows,'row_count':len(rows),'column_count':max(len(r) for r in rows)})
        return {'filename':name,'paragraphs':paragraphs,'tables':tables,'table_count':len(tables),'message':'DOCX tables are available for review and structured-table analysis.'}
    except Exception as exc: raise HTTPException(422,f'Unable to read DOCX: {exc}') from exc

@app.post('/api/v1/analysis/descriptive')
async def descriptive_endpoint(file:UploadFile=File(...),variables:str=''):
    df=load_dataframe(file.filename or 'upload.csv',await file.read()); selected=[x.strip() for x in variables.split(',') if x.strip()]
    if not selected: raise HTTPException(422,'Select at least one variable')
    try:return {'method':'descriptive','results':descriptive(df,selected)}
    except ValueError as exc: raise HTTPException(422,str(exc)) from exc

@app.post('/api/v1/analysis/association')
async def association_endpoint(file:UploadFile=File(...),variable:str='',outcome:str=''):
    df=load_dataframe(file.filename or 'upload.csv',await file.read())
    try:return {'method':'chi_square_fisher','results':categorical_association(df,variable,outcome)}
    except ValueError as exc: raise HTTPException(422,str(exc)) from exc

@app.post('/api/v1/analysis/diagnostic')
async def diagnostic_endpoint(file:UploadFile=File(...),index_test:str='',reference_standard:str=''):
    df=load_dataframe(file.filename or 'upload.csv',await file.read())
    try:return {'method':'diagnostic_accuracy','results':diagnostic_accuracy(df,index_test,reference_standard)}
    except ValueError as exc: raise HTTPException(422,str(exc)) from exc

@app.post('/api/v1/analysis/roc')
async def roc_endpoint(file:UploadFile=File(...),test:str='',outcome:str=''):
    df=load_dataframe(file.filename or 'upload.csv',await file.read())
    try:return {'method':'roc_auc','results':roc_auc(df,test,outcome)}
    except ValueError as exc: raise HTTPException(422,str(exc)) from exc

@app.post('/api/v1/analysis/dummy-table')
async def dummy_table_endpoint(dataset:UploadFile=File(...), specification:str=Form(...)):
    try:
        df=load_dataframe(dataset.filename or 'upload.csv',await dataset.read()); spec=json.loads(specification); return {'method':'dummy_table_analysis','results':analyse_dummy_table(df,spec)}
    except json.JSONDecodeError as exc: raise HTTPException(422,f'Invalid dummy-table specification JSON: {exc}') from exc
    except ValueError as exc: raise HTTPException(422,str(exc)) from exc

@app.post('/api/v1/dummy-table/analyze-template')
async def dummy_table_template_endpoint(dataset:UploadFile=File(...), dummy_table:UploadFile=File(...)):
    dataset_bytes=await dataset.read(); dummy_bytes=await dummy_table.read(); df=load_dataframe(dataset.filename or 'upload.csv',dataset_bytes)
    try:
        spec=build_spec(df,dummy_bytes,dummy_table.filename or 'dummy.xlsx'); return {'method':'dummy_table_template_analysis','results':{'dataset_filename':dataset.filename,'rows':int(df.shape[0]),'columns':int(df.shape[1]),'dummy_filename':dummy_table.filename,'sheets':spec['sheets'],'rows_spec':spec['rows_spec']}}
    except Exception as exc: raise HTTPException(422,f'Unable to parse dummy table: {exc}') from exc

@app.post('/api/v1/dummy-table/fill')
async def dummy_table_fill_endpoint(dataset:list[UploadFile]=File(...),dummy_table:UploadFile=File(...),outcome:str=Form(''),outcome_positive:str=Form(''),adjustment_variables:str=Form('')):
    try:
        dataset_bytes=[]; dataset_names=[]
        for upload in dataset: dataset_bytes.append(await upload.read()); dataset_names.append(upload.filename or 'upload.csv')
        dummy_bytes=await dummy_table.read(); dummy_name=dummy_table.filename or 'dummy.xlsx'; positive=[x.strip() for x in outcome_positive.split(',') if x.strip()] or None; adjustments=[x.strip() for x in adjustment_variables.split(',') if x.strip()] or None
        output,meta=fill_dummy_table(dataset_bytes,dataset_names,dummy_bytes,dummy_name,outcome=outcome.strip() or None,outcome_positive=positive,adjustment_variables=adjustments)
        base=dummy_name.rsplit('.',1)[0]; return StreamingResponse(io.BytesIO(output),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{base}_filled.xlsx"','X-Dummy-Analysis-Meta':json.dumps(meta,separators=(',',':'))})
    except ValueError as exc: raise HTTPException(422,str(exc)) from exc
    except Exception as exc: raise HTTPException(500,f'Unable to fill dummy table: {exc}') from exc

@app.post('/api/v1/analysis/rr-dummy')
async def rr_dummy_endpoint(file:UploadFile=File(...),config:str=Form(...)):
    try:
        cfg=json.loads(config); df=load_dataframe(file.filename or 'upload.csv',await file.read()); outcome_col=cfg['outcome']; exposure_col=cfg['exposure']; covariates=cfg.get('covariates',[])
        if outcome_col not in df.columns or exposure_col not in df.columns: raise HTTPException(422,'Outcome or exposure variable was not found in dataset')
        y=_binary(df[outcome_col],cfg.get('outcome_positive',[1,'1','yes','positive','unfavourable','no adherence'])); e=_binary(df[exposure_col],cfg.get('exposure_positive',[1,'1','intervention','case']))
        crude=crude_rr(y,e); adj=adjusted_rr(pd.DataFrame({'_y':y,'_e':e,**{c:df[c] for c in covariates if c in df.columns}}),'_y','_e',[c for c in covariates if c in df.columns],reference=0)
        return {'method':'crude_and_adjusted_rr','results':{'crude':crude,'adjusted':adj,'formatted':{'crude':format_rr(crude),'adjusted':format_rr(adj)}}}
    except json.JSONDecodeError as exc: raise HTTPException(422,f'Invalid RR configuration JSON: {exc}') from exc
    except KeyError as exc: raise HTTPException(422,f'Missing RR configuration field: {exc.args[0]}') from exc
    except ValueError as exc: raise HTTPException(422,str(exc)) from exc

@app.post('/api/v1/openeepi/calculate')
async def openeepi_calculate_endpoint(config:str=Form(...)):
    try:
        cfg=json.loads(config); module=cfg.pop('module'); result=openeepi_calculate(module,cfg); return {'method':f'OpenEpi:{module}','results':result}
    except KeyError as exc: raise HTTPException(422,f'Missing calculator field: {exc.args[0]}') from exc
    except (ValueError,TypeError,OverflowError,json.JSONDecodeError,IndexError) as exc: raise HTTPException(422,str(exc)) from exc

@app.post('/api/v1/openeepi/screening')
async def openeepi_screening_endpoint(config:str=Form(...)):
    try:
        cfg=json.loads(config); result=openeepi_screening(cfg['tp'],cfg['fp'],cfg['fn'],cfg['tn']); return {'method':'OpenEpi screening','results':result}
    except KeyError as exc: raise HTTPException(422,f'Missing screening cell: {exc.args[0]}') from exc
    except (ValueError,TypeError,json.JSONDecodeError) as exc: raise HTTPException(422,str(exc)) from exc

@app.post('/api/v1/analysis/poisson')
async def poisson_endpoint(file:UploadFile=File(...),config:str=Form(...)):
    try:
        cfg=json.loads(config); df=load_dataframe(file.filename or 'upload.csv',await file.read()); return {'method':'robust_poisson_rr','results':robust_poisson(df,cfg['outcome'],cfg.get('predictors',[]),cfg.get('categorical_predictors',[]),cfg.get('reference_categories',{}))}
    except KeyError as exc: raise HTTPException(422,f'Missing Poisson configuration field: {exc.args[0]}') from exc
    except (ValueError,np.linalg.LinAlgError,json.JSONDecodeError) as exc: raise HTTPException(422,str(exc)) from exc

@app.post('/api/v1/analysis/logistic')
async def logistic_endpoint(file:UploadFile=File(...),config:str=Form(...)):
    try:
        cfg=json.loads(config); df=load_dataframe(file.filename or 'upload.csv',await file.read()); return {'method':'logistic_regression','results':logistic_regression(df,cfg['outcome'],cfg.get('predictors',[]),cfg.get('categorical_predictors',[]),cfg.get('reference_categories',{}))}
    except KeyError as exc: raise HTTPException(422,f'Missing logistic configuration field: {exc.args[0]}') from exc
    except (ValueError,np.linalg.LinAlgError,json.JSONDecodeError) as exc: raise HTTPException(422,str(exc)) from exc

@app.post('/api/v1/analysis/table-one')
async def table_one_endpoint(file:UploadFile=File(...),variables:str='',group:str=''):
    df=load_dataframe(file.filename or 'upload.csv',await file.read()); selected=[x.strip() for x in variables.split(',') if x.strip()]
    if not selected: raise HTTPException(422,'Select at least one variable')
    try:return {'method':'table_one','results':table_one(df,selected,group.strip() or None)}
    except ValueError as exc: raise HTTPException(422,str(exc)) from exc

@app.post('/api/v1/data/quality')
async def quality_endpoint(file:UploadFile=File(...)):
    df=load_dataframe(file.filename or 'upload.csv',await file.read()); return {'method':'data_quality','results':data_quality(df)}
