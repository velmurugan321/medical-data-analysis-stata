"""Fill a supplied dummy Table 1 from one or more uploaded datasets.

The template is treated as a specification, not as a data source. Variable
headings and category rows are parsed from the workbook, dataset columns are
matched using explicit aliases, and results are written back into the original
workbook. No participant data are stored in the repository.

For the RR template the outcome is binary. Crude RR is calculated for each
category against the first category (reference). Adjusted RR uses Poisson
regression with robust variance and a configurable adjustment set. If no
adjustment set is supplied, adjusted RR is left blank rather than inventing a
model.
"""

import io
import re
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.stats import chi2_contingency, fisher_exact

from .rr_analysis import crude_rr, adjusted_rr, format_rr

SECTION_NAMES = {
    "demographic factors", "clinical factors", "clinical behavioural factors",
    "clinical / behavioural factors", "structural factors", "psychosocial factors",
}

VARIABLE_ALIASES = {
    "gender": ["sex", "gender"],
    "age": ["age", "age years", "age in years", "age group", "age group years"],
    "monthly income": ["monthly income", "income", "monthly household income", "monthly income inr"],
    "occupation": ["occupation", "job", "employment"],
    "phase of therapy": ["phase of therapy", "therapy phase", "treatment phase"],
    "category of tb": ["category of tb", "tb category", "previous treatment", "previously treated"],
    "type of tb": ["type of tb", "tb type", "disease type"],
    "transport mode to clinic": ["transport mode to clinic", "transport mode", "mode of transport", "transport"],
    "money spent to collect medication refills": ["money spent to collect medication refills", "money spent to collect refills", "money spent", "cost of refill", "refill cost"],
    "time spent to collect medication refills": ["time spent to collect medication refills", "time spent to collect refills", "time spent", "refill time", "time to collect medication"],
    "current tobacco use": ["current tobacco use", "tobacco use", "smoking", "tobacco"],
    "probable alcohol use": ["probable alcohol use", "alcohol use", "alcohol"],
}

OUTCOME_ALIASES = [
    "no adherence", "non adherence", "non-adherence", "urine test result",
    "programme outcome", "program outcome", "unfavourable", "unfavorable",
    "pib outcome", "perceived threat", "outcome",
]


def norm(v: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(v or "").lower()).strip()


def _canonical_variable(label):
    n = norm(label)
    for canonical, aliases in VARIABLE_ALIASES.items():
        if n == canonical or n in {norm(a) for a in aliases}:
            return canonical
    for canonical, aliases in VARIABLE_ALIASES.items():
        if any(norm(a) in n or n in norm(a) for a in aliases if norm(a)):
            return canonical
    return n


def _find_dataset_column(df, variable, categories=None):
    canonical = _canonical_variable(variable)
    aliases = VARIABLE_ALIASES.get(canonical, [variable])
    best = (None, 0.0)
    for col in df.columns:
        if str(col).startswith("__"):
            continue
        cn = norm(col)
        score = 0.0
        if cn == norm(variable) or cn == canonical:
            score = 1.0
        elif any(norm(a) == cn for a in aliases):
            score = 0.98
        elif any(norm(a) in cn or cn in norm(a) for a in aliases if norm(a)):
            score = 0.82
        if categories:
            vals = set(df[col].dropna().astype(str).map(norm).unique())
            hits = sum(1 for cat in categories if norm(cat) in vals)
            if hits:
                score = max(score, min(0.94, 0.55 + 0.35 * hits / len(categories)))
        if score > best[1]:
            best = (col, score)
    return best


def _find_outcome_column(df, requested=None):
    if requested and requested in df.columns:
        return requested
    if requested:
        n = norm(requested)
        for c in df.columns:
            if n == norm(c) or n in norm(c) or norm(c) in n:
                return c
    for alias in OUTCOME_ALIASES:
        n = norm(alias)
        for c in df.columns:
            cn = norm(c)
            if n == cn or n in cn or cn in n:
                return c
    return None


def _header_columns(ws):
    """Find output columns while tolerating merged/grouped Excel headers."""
    found = {}
    for r in range(1, min(ws.max_row, 12) + 1):
        for c in range(1, ws.max_column + 1):
            value = norm(ws.cell(r, c).value)
            if not value:
                continue
            if "variable category" in value or value == "variable":
                found.setdefault("variable", c)
            if "overall cohort" in value or value.startswith("overall"):
                found.setdefault("overall", c)
            if "no adherence" in value or "unfavourable" in value or "unfavorable" in value or "perceived threat" in value:
                found.setdefault("outcome", c)
            if value in {"p value", "pvalue", "p value"}:
                found.setdefault("p_value", c)
            if "unadjusted rr" in value or value == "unadjusted risk ratio":
                found.setdefault("unadjusted_rr", c)
            if value == "95% ci" or value == "95 ci" or "95% confidence" in value:
                # The first 95% CI belongs to unadjusted RR, the second to adjusted RR.
                if "unadjusted_ci" not in found:
                    found["unadjusted_ci"] = c
                elif "adjusted_ci" not in found:
                    found["adjusted_ci"] = c
            if "adjusted rr" in value or value == "adjusted risk ratio":
                found.setdefault("adjusted_rr", c)
    return found


def _site_arm_from_filename(name):
    n = norm(name)
    site = next((x for x in ("dindigul", "perambalur", "theni", "tiruvallur") if x in n), None)
    arm = "intervention" if "intervention" in n else ("control" if "control" in n else None)
    return site, arm


def _category_mask(series, category, variable):
    raw = series
    s = raw.astype(str).map(norm)
    c = norm(category)
    numeric = pd.to_numeric(raw, errors="coerce")
    canonical = _canonical_variable(variable)

    if canonical == "age":
        if re.search(r"18\s*29", c):
            return numeric.between(18, 29)
        if re.search(r"30\s*44", c):
            return numeric.between(30, 44)
        if "45" in c:
            return numeric >= 45

    if canonical == "monthly income":
        if "7500" in c and "14" not in c:
            return numeric < 7500
        if "7500" in c and "14" in c:
            return numeric.between(7500, 14999)
        if "15000" in c:
            return numeric >= 15000

    if canonical == "money spent to collect medication refills":
        if "0 24" in c:
            return numeric.between(0, 24)
        if "25 49" in c:
            return numeric.between(25, 49)
        if "50 75" in c:
            return numeric.between(50, 75)
        if "75" in c and ">" in c:
            return numeric > 75

    if canonical == "time spent to collect medication refills":
        if "30 minutes" in c and "60" not in c:
            return numeric < 30
        if "30 to 59" in c:
            return numeric.between(30, 59)
        if "60 to 239" in c:
            return numeric.between(60, 239)
        if "240" in c:
            return numeric >= 240

    exact = s == c
    if exact.any():
        return exact
    return s.str.contains(re.escape(c), regex=True, na=False)


def _test_p_from_groups(outcome, category_masks):
    # Overall variable p-value: outcome by the variable's mutually exclusive
    # categories. Missing observations are excluded.
    levels = []
    y = pd.to_numeric(outcome, errors="coerce")
    for label, mask in category_masks:
        valid = mask & y.notna()
        if valid.any():
            levels.append((label, valid))
    if len(levels) < 2:
        return None
    table = np.array([[int(((y == 1) & m).sum()), int(((y == 0) & m).sum())] for _, m in levels])
    if table.shape[0] < 2 or table.sum() == 0:
        return None
    try:
        if table.shape == (2, 2) and (table < 5).any():
            return float(fisher_exact(table)[1])
        return float(chi2_contingency(table, correction=False)[1])
    except Exception:
        return None


def _write(ws, row, col, value):
    if col:
        ws.cell(row, col).value = value


def fill_dummy_table(dataset_bytes_list, dataset_names, dummy_bytes, dummy_filename,
                     outcome=None, outcome_positive=None, adjustment_variables=None):
    frames = []
    for content, name in zip(dataset_bytes_list, dataset_names):
        ext = name.lower().rsplit(".", 1)[-1]
        if ext == "csv":
            df = pd.read_csv(io.BytesIO(content))
        elif ext == "tsv":
            df = pd.read_csv(io.BytesIO(content), sep="\t")
        elif ext in ("xlsx", "xls"):
            # Prefer a line-list-like sheet; otherwise use the first non-empty sheet.
            book = pd.ExcelFile(io.BytesIO(content))
            chosen = book.sheet_names[0]
            for s in book.sheet_names:
                probe = pd.read_excel(io.BytesIO(content), sheet_name=s, nrows=3)
                if probe.shape[1] >= 5 and probe.shape[0] > 0:
                    chosen = s
                    break
            df = pd.read_excel(io.BytesIO(content), sheet_name=chosen)
        elif ext == "dta":
            df = pd.read_stata(io.BytesIO(content))
        else:
            raise ValueError(f"Unsupported dataset type: {name}")
        df = df.copy()
        site, arm = _site_arm_from_filename(name)
        df["__source_dataset"] = name
        df["__source_site"] = site or ""
        df["__source_arm"] = arm or ""
        frames.append(df)
    if not frames:
        raise ValueError("At least one dataset is required")
    df = pd.concat(frames, ignore_index=True, sort=False)

    ext = dummy_filename.lower().rsplit(".", 1)[-1]
    if ext not in ("xlsx", "xls"):
        raise ValueError("Result filling currently requires an Excel dummy table (.xlsx/.xls)")
    wb = load_workbook(io.BytesIO(dummy_bytes))
    ws = wb[wb.sheetnames[0]]
    headers = _header_columns(ws)
    outcome_col = _find_outcome_column(df, outcome)
    if not outcome_col:
        raise ValueError("Could not identify the binary outcome variable. Select the outcome variable in the analysis page.")

    # Outcome coding: explicit values are preferred; otherwise common positive labels.
    if outcome_positive is None:
        outcome_positive = [1, "1", "yes", "positive", "unfavourable", "unfavorable", "no adherence", "non adherence", "non-adherence"]
    pos = {norm(x) for x in outcome_positive}
    y = df[outcome_col].map(lambda x: 1 if norm(x) in pos else (0 if pd.notna(x) else np.nan))
    if y.dropna().nunique() != 2:
        raise ValueError(f"Outcome '{outcome_col}' could not be converted to a binary 0/1 variable. Please specify outcome-positive values.")

    table_rows = []
    current_variable = None
    variable_keys = set(VARIABLE_ALIASES)
    for r in range(1, ws.max_row + 1):
        texts = [str(ws.cell(r, c).value).strip() for c in range(1, min(ws.max_column, 3) + 1)
                 if ws.cell(r, c).value is not None and str(ws.cell(r, c).value).strip()]
        if not texts:
            continue
        first = texts[0]
        nf = norm(first)
        if nf in SECTION_NAMES or nf in {"table 1 dummy template", "variable category variable", "variable category"}:
            continue
        can = _canonical_variable(first)
        if can in variable_keys:
            current_variable = first
            continue
        if current_variable:
            table_rows.append((r, current_variable, first))

    # Group category rows by variable for variable-level p-values and reference selection.
    grouped = {}
    for r, variable, category in table_rows:
        grouped.setdefault(variable, []).append((r, category))

    adjustment_variables = adjustment_variables or []
    adjustment_cols = []
    for av in adjustment_variables:
        col, score = _find_dataset_column(df, av)
        if col and score >= 0.70 and col not in adjustment_cols and col != outcome_col:
            adjustment_cols.append(col)

    filled = 0
    unmapped = []
    for variable, rows in grouped.items():
        categories = [x[1] for x in rows]
        col, confidence = _find_dataset_column(df, variable, categories)
        if col is None or confidence < 0.50:
            for r, category in rows:
                unmapped.append({"row": r, "variable": variable, "category": category, "reason": "dataset variable not confidently mapped"})
            continue

        masks = [(cat, _category_mask(df[col], cat, variable)) for _, cat in rows]
        p = _test_p_from_groups(y, masks)
        reference_category = categories[0] if categories else None
        reference_mask = _category_mask(df[col], reference_category, variable) if reference_category else None

        for (r, category), (_, mask) in zip(rows, masks):
            valid = mask & y.notna()
            n = int(valid.sum())
            positive_n = int(((y == 1) & valid).sum())
            pct = 100 * positive_n / n if n else None
            _write(ws, r, headers.get("overall"), str(n) if n else "")
            _write(ws, r, headers.get("outcome"), f"{positive_n} ({pct:.1f}%)" if pct is not None else "")
            _write(ws, r, headers.get("p_value"), _format_p(p))
            filled += 1

            # Crude RR: category vs first category (reference).
            if reference_mask is not None and category != reference_category:
                exposure = pd.Series(np.nan, index=df.index)
                exposure[mask] = 1
                exposure[reference_mask] = 0
                crude = crude_rr(y, exposure)
                if crude:
                    _write(ws, r, headers.get("unadjusted_rr"), f"{crude['rr']:.2f}")
                    _write(ws, r, headers.get("unadjusted_ci"), f"{crude['ci_low']:.2f}–{crude['ci_high']:.2f}")
                    filled += 2

                    if adjustment_cols:
                        model_df = df[[outcome_col, col] + adjustment_cols].copy()
                        model_df["_y"] = y
                        model_df["_e"] = exposure
                        covs = [c for c in adjustment_cols if c != col]
                        # Use the already binary exposure for the target category.
                        adj = adjusted_rr(model_df, "_y", "_e", covs, reference=0)
                        if adj and not adj.get("error"):
                            _write(ws, r, headers.get("adjusted_rr"), f"{adj['rr']:.2f}")
                            _write(ws, r, headers.get("adjusted_ci"), f"{adj['ci_low']:.2f}–{adj['ci_high']:.2f}")
                            filled += 2

            else:
                _write(ws, r, headers.get("unadjusted_rr"), "Ref")
                _write(ws, r, headers.get("unadjusted_ci"), "Ref")

        # If the template has a variable-heading row immediately before the
        # categories, put the variable-level p-value there as well.
        first_row = rows[0][0] - 1
        if first_row > 0 and headers.get("p_value") and ws.cell(first_row, 1).value:
            if _canonical_variable(ws.cell(first_row, 1).value) == _canonical_variable(variable):
                _write(ws, first_row, headers["p_value"], _format_p(p))

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue(), {
        "rows_in_dataset": int(df.shape[0]),
        "columns_in_dataset": int(len([c for c in df.columns if not str(c).startswith("__")])),
        "datasets": dataset_names,
        "outcome_variable": outcome_col,
        "outcome_positive_values": list(outcome_positive),
        "adjustment_variables": adjustment_cols,
        "filled_cells": filled,
        "parsed_table_rows": len(table_rows),
        "unmapped_rows": unmapped[:100],
        "site_arm_sources": [{"dataset": n, "site": _site_arm_from_filename(n)[0], "arm": _site_arm_from_filename(n)[1]} for n in dataset_names],
        "rr_method": "crude 2x2 risk ratio; adjusted Poisson regression with robust variance",
        "reference_rule": "first category listed in each dummy-table variable",
    }


def _format_p(p):
    if p is None or not np.isfinite(p):
        return ""
    return "<0.001" if p < 0.001 else f"{p:.3f}"
